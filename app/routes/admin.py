# app/routes/admin.py
from flask import Blueprint, request, redirect, url_for, render_template, jsonify, session
from flask_login import login_required, current_user, login_user
from app.db import db
from app.models import User, Device, Order

bp = Blueprint("admin", __name__)

# ---- Политика доступа ----

def require_superadmin():
    return current_user.is_authenticated and current_user.role == "superadmin"

def forbid():
    return "403 Forbidden", 403

# ---- Меню суперадмина ----
@bp.route("/")
@login_required
def dashboard():
    if not require_superadmin():
        return forbid()
    return render_template("admin/dashboard.html")

# ---- Отдельные страницы ----
@bp.route("/users")
@login_required
def users_page():
    if not require_superadmin():
        return forbid()
    return render_template("admin/users_list.html")

@bp.route("/users/new")
@login_required
def users_new_page():
    if not require_superadmin():
        return forbid()
    return render_template("admin/users_add.html")

# ---- Analytics pages ----
@bp.route("/analytics")
@login_required
def analytics_page():
    if not require_superadmin():
        return forbid()
    return render_template("admin/analytics.html")

# ==================== USERS API ====================
@bp.route("/api/users", methods=["GET"])
@login_required
def api_users_list():
    if not require_superadmin():
        return forbid()
    users = User.query.order_by(User.id.asc()).all()
    data = [{"id": u.id, "email": u.email, "role": u.role, "is_enabled": getattr(u, "is_enabled", True)} for u in users]
    return jsonify(data)

@bp.route("/api/users", methods=["POST"])
@login_required
def api_users_add():
    if not require_superadmin():
        return forbid()
    payload = request.get_json(force=True)
    email = (payload.get("email") or "").strip()
    password = (payload.get("password") or "").strip()
    role = payload.get("role", "localadmin")
    if not email or not password or role not in ("localadmin", "superadmin"):
        return jsonify({"error": "bad input"}), 400
    if User.query.filter_by(email=email).first():
        return jsonify({"error": "exists"}), 409
    u = User(email=email, role=role)
    u.set_password(password)
    db.session.add(u)
    db.session.commit()
    return jsonify({"ok": True, "id": u.id})

@bp.route("/api/users/<int:uid>", methods=["DELETE"])
@login_required
def api_users_delete(uid):
    if not require_superadmin():
        return forbid()
    if uid == current_user.id:
        return jsonify({"error":"cannot delete yourself"}), 400
    u = User.query.get(uid)
    if not u:
        return jsonify({"error": "not found"}), 404
    if u.role == "superadmin":
        return jsonify({"error":"cannot delete superadmin"}), 400
    db.session.delete(u)
    db.session.commit()
    return jsonify({"ok": True})

# ==================== ANALYTICS API ====================
@bp.route("/api/analytics/summary", methods=["GET"])
@login_required
def api_analytics_summary():
    if not require_superadmin():
        return forbid()
    from datetime import datetime, timedelta
    range_ = request.args.get("range", "month")
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")
    now = datetime.utcnow()
    start = None
    end = None
    if date_from or date_to:
        if date_from:
            try:
                start = datetime.strptime(date_from, "%Y-%m-%d")
            except ValueError:
                start = None
        if date_to:
            try:
                end = datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)
            except ValueError:
                end = None
    else:
        if range_ == "day":
            start = now - timedelta(days=1)
        elif range_ == "week":
            start = now - timedelta(days=7)
        elif range_ == "month":
            start = now - timedelta(days=30)

    base_query = Order.query
    if start:
        base_query = base_query.filter(Order.created_at >= start)
    if end:
        base_query = base_query.filter(Order.created_at < end)

    total_orders = base_query.with_entities(db.func.count(Order.id)).scalar() or 0
    total_amount = base_query.with_entities(db.func.coalesce(db.func.sum(Order.amount), 0)).scalar() or 0
    total_minutes = base_query.with_entities(db.func.coalesce(db.func.sum(Order.minutes), 0)).scalar() or 0

    devices = Device.query.order_by(Device.id.asc()).all()
    by_device = []
    for d in devices:
        base = db.session.query(Order).filter(Order.device_id == (d.device_uid or str(d.id)))
        if start:
            base = base.filter(Order.created_at >= start)
        if end:
            base = base.filter(Order.created_at < end)
        cnt = base.with_entities(db.func.count(Order.id)).scalar() or 0
        amt = base.with_entities(db.func.coalesce(db.func.sum(Order.amount), 0)).scalar() or 0
        # Конвертируем копейки в рубли (делим на 100)
        by_device.append({"id": d.id, "name": d.name, "orders": int(cnt), "amount": round(amt / 100, 2)})
    return jsonify({
        "total_orders": int(total_orders),
        "total_amount": round(total_amount / 100, 2),  # Конвертируем копейки в рубли
        "total_minutes": int(total_minutes or 0),
        "by_device": by_device,
    })

@bp.route("/api/analytics/export.xlsx", methods=["GET"])
@login_required
def api_analytics_export():
    if not require_superadmin():
        return forbid()
    
    from io import BytesIO
    from datetime import datetime, timedelta
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception:
        return jsonify({"error": "xlsx_not_supported"}), 500
    
    # Получаем параметры фильтрации
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")
    
    # Строим запрос
    query = Order.query
    
    if date_from:
        try:
            date_from_obj = datetime.strptime(date_from, "%Y-%m-%d")
            query = query.filter(Order.created_at >= date_from_obj)
        except ValueError:
            pass
    
    if date_to:
        try:
            # Добавляем 1 день, чтобы включить весь день
            date_to_obj = datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(Order.created_at < date_to_obj)
        except ValueError:
            pass
    
    orders = query.order_by(Order.created_at.asc()).all()
    
    # Создаем Excel файл
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    
    # Заголовки с форматированием
    headers = ["ID", "Дата (ДД.ММ.ГГГГ)", "Время (ЧЧ:ММ:СС)", "Сумма (₽)", "Валюта", "Device ID", "Минуты", "Статус"]
    ws.append(headers)
    
    # Форматирование заголовков
    header_fill = PatternFill(start_color="5B8CFF", end_color="5B8CFF", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Заполняем данные
    total_orders = 0
    total_minutes = 0
    total_amount_kop = 0
    total_time_seconds = 0
    total_time_count = 0
    for o in orders:
        created_at = o.created_at
        date_str = created_at.strftime("%d.%m.%Y") if created_at else ""
        time_str = created_at.strftime("%H:%M:%S") if created_at else ""
        amount_rub = round(o.amount / 100, 2) if o.amount else 0

        total_orders += 1
        if o.minutes:
            total_minutes += int(o.minutes)
        if o.amount:
            total_amount_kop += int(o.amount)
        if created_at:
            total_time_seconds += created_at.hour * 3600 + created_at.minute * 60 + created_at.second
            total_time_count += 1
        
        ws.append([
            o.id,
            date_str,
            time_str,
            amount_rub,
            o.currency or "RUB",
            o.device_id or "",
            o.minutes or 0,
            o.payment_status or "unknown"
        ])

    total_amount_rub = round(total_amount_kop / 100, 2)
    if total_time_count:
        avg_seconds = int(round(total_time_seconds / total_time_count))
        avg_time_str = f"{avg_seconds // 3600:02d}:{(avg_seconds % 3600) // 60:02d}:{avg_seconds % 60:02d}"
    else:
        avg_time_str = "—"
    total_row = [
        "Итого",
        f"Кол-во операций: {total_orders}",
        f"Среднее время: {avg_time_str}",
        f"{total_amount_rub:.2f} RUB",
        "",
        "",
        f"{total_minutes} Минут",
        ""
    ]
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Автоширина колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    from flask import send_file
    filename = f"analytics_{date_from or 'all'}_{date_to or 'current'}.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@bp.route("/api/users/<int:uid>/toggle", methods=["POST"])
@login_required
def api_users_toggle(uid):
    if not require_superadmin():
        return forbid()
    u = User.query.get(uid)
    if not u:
        return jsonify({"error":"not found"}), 404
    if u.role != "localadmin":
        return jsonify({"error":"only localadmin can be toggled"}), 400
    enabled = not getattr(u, "is_enabled", True)
    # если у модели ещё нет поля (на старой БД) — считаем включённым
    try:
        u.is_enabled = enabled
        from app.models import User as U
        U.query.filter_by(parent_id=u.id).update({"is_enabled": True if enabled else False})
        db.session.commit()
    except Exception:
        pass
    Device.query.filter_by(owner_id=u.id).update({"is_active": True if enabled else False})
    db.session.commit()
    return jsonify({"ok": True, "is_enabled": enabled})

@bp.route("/impersonate/<int:uid>", methods=["POST"])
@login_required
def impersonate(uid):
    if not require_superadmin():
        return forbid()
    target = User.query.get(uid)
    if not target or target.role != "localadmin":
        return jsonify({"error":"not found"}), 404
    if getattr(target, "is_enabled", True) is False:
        return jsonify({"error":"admin is disabled"}), 400
    session["sudo_original_id"] = current_user.id
    login_user(target, force=True)
    return jsonify({"ok": True, "redirect": url_for("local_admin.dashboard")})

@bp.route("/unsudo", methods=["POST"])
@login_required
def unsudo():
    orig_id = session.pop("sudo_original_id", None)
    if not orig_id:
        return jsonify({"ok": False}), 400
    admin = User.query.get(orig_id)
    if not admin or admin.role != "superadmin":
        return jsonify({"ok": False}), 400
    login_user(admin, force=True)
    return jsonify({"ok": True, "redirect": url_for("admin.dashboard")})
