# app/routes/local_admin.py
from datetime import datetime, timedelta, timezone
import re
from flask import Blueprint, abort, render_template, request, jsonify, session, send_file, url_for, current_app
from flask_login import login_required, current_user
from app.db import db
from app.models import Device, User, Order, ClientConsent, schedule_activation, clear_expired_activation
from sqlalchemy.exc import IntegrityError

bp = Blueprint("local_admin", __name__)

def is_local():  return current_user.is_authenticated and current_user.role == "localadmin"
def is_worker(): return current_user.is_authenticated and current_user.role == "worker"
def can_control(): return current_user.is_authenticated and current_user.role in ("localadmin", "worker")
def owner_id_for_user():
    if is_local():
        return current_user.id
    if is_worker():
        return current_user.parent_id
    return None

_TZ_OFFSET_RE = re.compile(r"^([+-])(\d{2})(?::?(\d{2}))?$")

def parse_tz_offset(raw_value: str):
    if not raw_value:
        return timezone.utc
    value = str(raw_value).strip().upper()
    if value in ("UTC", "GMT", "Z", "+00:00", "+00", "0", "00:00", "-00:00", "-00"):
        return timezone.utc
    match = _TZ_OFFSET_RE.match(value)
    if not match:
        return timezone.utc
    sign, hh, mm = match.groups()
    try:
        hours = int(hh)
        minutes = int(mm or 0)
    except ValueError:
        return timezone.utc
    if hours > 14 or minutes >= 60:
        return timezone.utc
    total = hours * 60 + minutes
    if sign == "-":
        total = -total
    return timezone(timedelta(minutes=total))

# -------- Pages --------
@bp.route("/")
@login_required
def dashboard():
    if not is_local():
        abort(403)
    is_sudo = bool(session.get("sudo_original_id"))
    return render_template("local/dashboard.html", is_sudo=is_sudo, email=current_user.email)

@bp.route("/devices")
@login_required
def devices_page():
    if not is_local():
        abort(403)
    return render_template("local/devices_list.html")

@bp.route("/devices/new")
@login_required
def devices_new_page():
    if not is_local():
        abort(403)
    return render_template("local/devices_new.html")

@bp.route("/devices/<int:did>")
@login_required
def devices_edit_page(did):
    if not is_local():
        abort(403)
    return render_template("local/devices_edit.html", did=did)

@bp.route("/worker")
@login_required
def worker_page():
    if not (is_local() or is_worker()):
        abort(403)
    return render_template("local/worker.html")


@bp.route("/worker/device/<int:did>")
@login_required
def worker_device_page(did):
    if not (is_local() or is_worker()):
        abort(403)
    return render_template("local/worker_device.html", did=did)

@bp.route("/stats")
@login_required
def stats_page():
    if not (is_local() or is_worker()):
        abort(403)
    return render_template("local/stats.html")

# -------- Local analytics API --------
@bp.route("/api/analytics/summary", methods=["GET"])
@login_required
def api_local_analytics_summary():
    if not is_local():
        abort(403)
    # only devices belonging to current local admin
    devs = Device.query.filter_by(owner_id=current_user.id).all()
    ids = [d.device_uid or str(d.id) for d in devs]
    total_orders = db.session.query(db.func.count(Order.id)).filter(Order.device_id.in_(ids)).scalar() or 0
    total_amount = db.session.query(db.func.coalesce(db.func.sum(Order.amount), 0)).filter(Order.device_id.in_(ids)).scalar() or 0
    total_minutes = db.session.query(db.func.coalesce(db.func.sum(Order.minutes), 0)).filter(Order.device_id.in_(ids)).scalar() or 0
    by_device = []
    for d in devs:
        did = d.device_uid or str(d.id)
        cnt = db.session.query(db.func.count(Order.id)).filter(Order.device_id == did).scalar() or 0
        amt = db.session.query(db.func.coalesce(db.func.sum(Order.amount), 0)).filter(Order.device_id == did).scalar() or 0
        # Конвертируем копейки в рубли
        by_device.append({"id": d.id, "name": d.name, "orders": int(cnt), "amount": round(amt / 100, 2)})
    return jsonify({
        "total_orders": int(total_orders),
        "total_amount": round(total_amount / 100, 2),
        "total_minutes": int(total_minutes or 0),
        "by_device": by_device
    })

@bp.route("/api/analytics/export.xlsx", methods=["GET"])
@login_required
def api_local_analytics_export():
    if not is_local():
        abort(403)
    
    from io import BytesIO
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception:
        return jsonify({"error": "xlsx_not_supported"}), 500
    
    # Получаем параметры фильтрации
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")
    tz_offset = request.args.get("tz_offset")
    tz = parse_tz_offset(tz_offset)
    
    devs = Device.query.filter_by(owner_id=current_user.id).all()
    ids = [d.device_uid or str(d.id) for d in devs]
    
    # Строим запрос
    query = Order.query.filter(Order.device_id.in_(ids))
    
    if date_from:
        try:
            date_from_obj = datetime.strptime(date_from, "%Y-%m-%d").date()
            start_local = datetime.combine(date_from_obj, datetime.min.time()).replace(tzinfo=tz)
            start_utc = start_local.astimezone(timezone.utc).replace(tzinfo=None)
            query = query.filter(Order.created_at >= start_utc)
        except ValueError:
            pass
    
    if date_to:
        try:
            date_to_obj = datetime.strptime(date_to, "%Y-%m-%d").date() + timedelta(days=1)
            end_local = datetime.combine(date_to_obj, datetime.min.time()).replace(tzinfo=tz)
            end_utc = end_local.astimezone(timezone.utc).replace(tzinfo=None)
            query = query.filter(Order.created_at < end_utc)
        except ValueError:
            pass
    
    orders = query.order_by(Order.created_at.asc()).all()
    
    # Создаем Excel файл
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    
    # Заголовки
    headers = ["ID", "Дата (ДД.ММ.ГГГГ)", "Время (ЧЧ:ММ:СС)", "Сумма (₽)", "Валюта", "Device ID", "Минуты", "Статус"]
    ws.append(headers)
    
    # Форматирование заголовков
    header_fill = PatternFill(start_color="5B8CFF", end_color="5B8CFF", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Заполняем данные
    total_orders = 0
    total_minutes = 0
    total_amount_kop = 0
    total_time_seconds = 0
    total_time_count = 0
    for o in orders:
        created_at = o.created_at
        if created_at:
            local_dt = created_at.replace(tzinfo=timezone.utc).astimezone(tz)
            date_str = local_dt.strftime("%d.%m.%Y")
            time_str = local_dt.strftime("%H:%M:%S")
        else:
            date_str = ""
            time_str = ""
        amount_rub = round(o.amount / 100, 2) if o.amount else 0

        total_orders += 1
        if o.minutes:
            total_minutes += int(o.minutes)
        if o.amount:
            total_amount_kop += int(o.amount)
        if created_at:
            total_time_seconds += local_dt.hour * 3600 + local_dt.minute * 60 + local_dt.second
            total_time_count += 1
        
        ws.append([
            o.id,
            date_str,
            time_str,
            amount_rub,
            o.currency or "RUB",
            o.device_id or "",
            o.minutes or 0,
            o.payment_status or "unknown"
        ])

    total_amount_rub = round(total_amount_kop / 100, 2)
    if total_time_count:
        avg_seconds = int(round(total_time_seconds / total_time_count))
        avg_time_str = f"{avg_seconds // 3600:02d}:{(avg_seconds % 3600) // 60:02d}:{avg_seconds % 60:02d}"
    else:
        avg_time_str = "—"
    total_row = [
        "Итого",
        f"Кол-во операций: {total_orders}",
        f"Среднее время: {avg_time_str}",
        f"{total_amount_rub:.2f} RUB",
        "",
        "",
        f"{total_minutes} Минут",
        ""
    ]
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Автоширина колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    from flask import send_file
    filename = f"analytics_local_{date_from or 'all'}_{date_to or 'current'}.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@bp.route("/workers")
@login_required
def workers_page():
    if not is_local():
        abort(403)
    return render_template("local/workers_list.html")

@bp.route("/workers/new")
@login_required
def workers_new_page():
    if not is_local():
        abort(403)
    return render_template("local/workers_add.html")

# -------- Workers: API --------
@bp.route("/api/workers", methods=["GET"])
@login_required
def api_workers_list():
    if not is_local():
        abort(403)
    rows = User.query.filter_by(parent_id=current_user.id, role="worker").order_by(User.id.asc()).all()
    return jsonify([{"id": u.id, "email": u.email, "is_enabled": u.is_enabled, "password": u.worker_password} for u in rows])

@bp.route("/api/workers", methods=["POST"])
@login_required
def api_workers_add():
    if not is_local():
        abort(403)
    p = request.get_json(force=True)
    email = (p.get("email") or "").strip()
    password = (p.get("password") or "").strip()
    if not email or not password:
        return jsonify({"error": "bad input"}), 400
    if User.query.filter_by(email=email).first():
        return jsonify({"error": "exists"}), 409
    u = User(email=email, role="worker", parent_id=current_user.id, is_enabled=True)
    u.set_password(password)
    u.worker_password = password
    db.session.add(u); db.session.commit()
    return jsonify({"ok": True, "id": u.id})

@bp.route("/api/workers/<int:wid>/toggle", methods=["POST"])
@login_required
def api_workers_toggle(wid):
    if not is_local():
        abort(403)
    u = User.query.filter_by(id=wid, parent_id=current_user.id, role="worker").first()
    if not u:
        return jsonify({"error": "not found"}), 404
    u.is_enabled = not u.is_enabled
    db.session.commit()
    return jsonify({"ok": True, "is_enabled": u.is_enabled})

@bp.route("/api/workers/<int:wid>", methods=["DELETE"])
@login_required
def api_workers_delete(wid):
    if not is_local():
        abort(403)
    u = User.query.filter_by(id=wid, parent_id=current_user.id, role="worker").first()
    if not u:
        return jsonify({"error": "not found"}), 404
    db.session.delete(u)
    db.session.commit()
    return jsonify({"ok": True})

# -------- Devices CRUD (локальный только) --------
@bp.route("/api/devices", methods=["GET"])
@login_required
def api_devices_list():
    if not is_local():
        abort(403)
    from app.models import is_online
    rows = Device.query.filter_by(owner_id=current_user.id).order_by(Device.id.asc()).all()
    data = []
    for d in rows:
        data.append(dict(
            id=d.id,
            name=d.name,
            device_uid=d.device_uid,
            price_per_minute=d.price_per_minute,
            is_active=d.is_active,
            relay_state=d.relay_state,
            active_until=d.active_until.isoformat()+"Z" if d.active_until else None,
            last_seen_at=d.last_seen_at.isoformat()+"Z" if d.last_seen_at else None,
            online=is_online(d),
            allowed_minutes=d.allowed_minutes or None
        ))
    return jsonify(data)


@bp.route("/api/devices", methods=["POST"])
@login_required
def api_devices_add():
    if not is_local():
        abort(403)
    p = request.get_json(force=True)
    name = (p.get("name") or "").strip()
    uid = (p.get("device_uid") or "").strip()
    price = int(p.get("price_per_minute") or 0)
    allowed = p.get("allowed_minutes")
    if isinstance(allowed, list):
        try:
            allowed = sorted({int(x) for x in allowed if 1 <= int(x) <= 720})
        except Exception:
            allowed = None
    else:
        allowed = None

    if not name or not uid:
        return jsonify({"error": "Заполните имя и UID"}), 400

    if Device.query.filter_by(device_uid=uid).first():
        return jsonify({"error": "UID уже существует"}), 409

    d = Device(
        name=name, device_uid=uid, price_per_minute=price,
        owner_id=current_user.id, is_active=True, allowed_minutes=allowed
    )
    db.session.add(d)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({"error": "UID уже существует"}), 409

    return jsonify({"ok": True, "id": d.id})


@bp.route("/api/devices/<int:did>", methods=["GET"])
@login_required
def api_device_get(did):
    if not is_local():
        abort(403)
    from app.models import is_online
    d = Device.query.filter_by(id=did, owner_id=current_user.id).first()
    if not d:
        return jsonify({"error": "not found"}), 404
    return jsonify(dict(
        id=d.id,
        name=d.name,
        device_uid=d.device_uid,
        price_per_minute=d.price_per_minute,
        is_active=d.is_active,
        relay_state=d.relay_state,
        active_until=d.active_until.isoformat()+"Z" if d.active_until else None,
        last_seen_at=d.last_seen_at.isoformat()+"Z" if d.last_seen_at else None,
        online=is_online(d),
        allowed_minutes=d.allowed_minutes or None
    ))


@bp.route("/api/devices/<int:did>", methods=["PUT"])
@login_required
def api_device_update(did):
    if not is_local():
        abort(403)
    d = Device.query.filter_by(id=did, owner_id=current_user.id).first()
    if not d:
        return jsonify({"error": "not found"}), 404
    p = request.get_json(force=True)
    for k in ("name", "device_uid", "price_per_minute", "is_active"):
        if k in p:
            setattr(d, k, p[k])
    if "allowed_minutes" in p:
        allowed = p.get("allowed_minutes")
        if isinstance(allowed, list):
            try:
                allowed = sorted({int(x) for x in allowed if 1 <= int(x) <= 720})
            except Exception:
                allowed = None
        else:
            allowed = None
        d.allowed_minutes = allowed
    db.session.commit()
    return jsonify({"ok": True})

@bp.route("/api/devices/<int:did>", methods=["DELETE"])
@login_required
def api_device_delete(did):
    if not is_local():
        abort(403)
    d = Device.query.filter_by(id=did, owner_id=current_user.id).first()
    if not d:
        return jsonify({"error": "not found"}), 404
    db.session.delete(d)
    db.session.commit()
    return jsonify({"ok": True})

@bp.route("/api/devices/<int:did>/toggle", methods=["POST"])
@login_required
def api_device_toggle(did):
    if not is_local():
        abort(403)
    d = Device.query.filter_by(id=did, owner_id=current_user.id).first()
    if not d:
        return jsonify({"error": "not found"}), 404
    d.is_active = not d.is_active
    db.session.commit()
    return jsonify({"ok": True, "is_active": d.is_active})

# -------- Worker status / bulk --------
@bp.route("/api/worker/status")
@login_required
def api_worker_status():
    if current_user.role not in ("localadmin", "worker"):
        return jsonify({"error": "forbidden"}), 403

    owner_id = owner_id_for_user()

    if not owner_id:
        return jsonify([])

    now = datetime.utcnow()
    from app.models import is_online
    rows = Device.query.filter_by(owner_id=owner_id).all()

    out = []
    for d in rows:
        remaining = 0
        starts_in = 0
        if d.relay_state and d.active_until:
            if d.active_until > now:
                remaining = int((d.active_until - now).total_seconds())
                if d.relay_start_at and d.relay_start_at > now:
                    starts_in = int((d.relay_start_at - now).total_seconds())
            else:
                clear_expired_activation(d, now=now)
        out.append({
            "id": d.id,
            "uid": d.device_uid,
            "name": d.name,
            "is_active": d.is_active,
            "relay_state": d.relay_state,
            "remaining_seconds": remaining,
            "starts_in_seconds": starts_in,
            "last_seen_at": d.last_seen_at.isoformat()+"Z" if d.last_seen_at else None,
            "online": is_online(d)
        })

    db.session.commit()
    return jsonify(out)


@bp.route("/api/worker/device/<int:did>", methods=["GET"])
@login_required
def api_worker_device(did):
    if not can_control():
        return jsonify({"error": "forbidden"}), 403

    owner_id = owner_id_for_user()
    if not owner_id:
        return jsonify({"error": "forbidden"}), 403

    now = datetime.utcnow()
    from app.models import is_online
    d = Device.query.filter_by(id=did, owner_id=owner_id).first()
    if not d:
        return jsonify({"error": "not found"}), 404

    remaining = 0
    starts_in = 0
    if d.relay_state and d.active_until:
        if d.active_until > now:
            remaining = int((d.active_until - now).total_seconds())
            if d.relay_start_at and d.relay_start_at > now:
                starts_in = int((d.relay_start_at - now).total_seconds())
        else:
            clear_expired_activation(d, now=now)
            db.session.commit()

    return jsonify({
        "id": d.id,
        "uid": d.device_uid,
        "name": d.name,
        "is_active": d.is_active,
        "relay_state": d.relay_state,
        "remaining_seconds": remaining,
        "starts_in_seconds": starts_in,
        "online": is_online(d),
        "allowed_minutes": d.allowed_minutes or None,
    })


@bp.route("/api/worker/device/<int:did>/start", methods=["POST"])
@login_required
def api_worker_device_start(did):
    if not can_control():
        return jsonify({"error": "forbidden"}), 403

    owner_id = owner_id_for_user()
    if not owner_id:
        return jsonify({"error": "forbidden"}), 403

    payload = request.get_json(force=True) or {}
    minutes = int(payload.get("minutes") or 0)
    if minutes <= 0 or minutes > 720:
        return jsonify({"error": "bad_input"}), 400

    now = datetime.utcnow()
    from app.models import is_online
    d = Device.query.filter_by(id=did, owner_id=owner_id).first()
    if not d or not d.is_active:
        return jsonify({"error": "not found"}), 404
    if not is_online(d):
        return jsonify({"error": "device_offline"}), 409

    if d.relay_state and d.active_until and d.active_until > now:
        remaining = int((d.active_until - now).total_seconds())
        return jsonify({"error": "busy", "remaining_seconds": remaining}), 409

    if d.allowed_minutes:
        try:
            normalize = sorted({int(x) for x in d.allowed_minutes})
        except Exception:
            normalize = None
        if normalize and minutes not in normalize:
            return jsonify({"error": "minutes_not_allowed"}), 400

    delay_seconds = int(current_app.config.get("RELAY_DELAY_SECONDS", 0) or 0)
    schedule_activation(d, minutes, now=now, delay_seconds=delay_seconds)

    amount = int(d.price_per_minute or 0) * minutes
    db.session.add(Order(
        amount=amount,
        currency="RUB",
        device_id=d.device_uid or str(d.id),
        device_table_id=d.id,
        minutes=minutes,
        payment_status="succeeded"
    ))

    db.session.commit()
    return jsonify({"ok": True, "until": d.active_until.isoformat() + "Z"})


@bp.route("/api/worker/device/<int:did>/stop", methods=["POST"])
@login_required
def api_worker_device_stop(did):
    if not can_control():
        return jsonify({"error": "forbidden"}), 403

    owner_id = owner_id_for_user()
    if not owner_id:
        return jsonify({"error": "forbidden"}), 403

    d = Device.query.filter_by(id=did, owner_id=owner_id).first()
    if not d:
        return jsonify({"error": "not found"}), 404

    d.relay_state = False
    d.active_until = None
    d.relay_start_at = None
    db.session.commit()
    return jsonify({"ok": True})


@bp.route("/consents")
@login_required
def consents_page():
    if not is_local():
        abort(403)
    return render_template("local/consents.html")


@bp.route("/api/consents", methods=["GET"])
@login_required
def api_consents_list():
    if not is_local():
        abort(403)

    rows = ClientConsent.query.filter_by(owner_id=current_user.id).order_by(ClientConsent.id.desc()).limit(500).all()
    data = []
    for row in rows:
        data.append({
            "id": row.id,
            "date": row.consent_date.isoformat() if row.consent_date else "",
            "time": row.consent_time.strftime("%H:%M:%S") if row.consent_time else "",
            "email": row.email,
            "agreed": bool(row.agreed),
        })
    return jsonify(data)

# bulk power endpoint removed by requirement

@bp.route("/devices/<int:did>/qr", methods=["GET"])
@login_required
def device_qr(did):
    # только локальный админ
    if current_user.role != "localadmin":
        abort(403)

    d = Device.query.filter_by(id=did, owner_id=current_user.id).first()
    if not d:
        abort(404)

    # Абсолютная ссылка на публичную страницу устройства
    # endpoint: public.public_device (bp 'public', func 'public_device')
    link = url_for("public.public_device", uid=d.device_uid, _external=True)

    # Генерируем PNG QR
    import qrcode
    from io import BytesIO
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=2,
    )
    qr.add_data(link)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buf = BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    fname = f"device-{d.device_uid}.png"
    return send_file(buf, mimetype="image/png", as_attachment=True, download_name=fname)

